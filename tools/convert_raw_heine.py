#!/usr/bin/env python3
"""Heine et al. 2025 の生データ（Nils Heine 提供、2026-10-13）を推定用の縦長 CSV にする。

入力（data_5species/experiment_data/raw_heine/）:
  qpcr_species_distribution.xlsx  — qPCR の種の割合（%）。シート Commensal / Dysbiotic、
      ブロック「<Static|HOBIC> all cells」（PMA なし＝全菌）と「<…> only living cells」（PMA あり＝生きた菌）、
      行 = 日（Tag）、列 = 種ごとの測定（反復 × 技術反復）。
  livedead_HOBIC_Static.xlsx      — LIVE/DEAD（SYTO9/PI）の画像ごとの値。シート HOBIC22・Static22（commensal）、
      HOBIC24・Static24（dysbiotic）。左: Biofilmvolumen（日ごと）、右: Lebend / Tot（%）を日ごとに 2 列。
      HOBIC22 の Day 1 は 40 倍対物で撮影 → 63 倍に換算した体積が下の別ブロック（"Umgerechnet auf 63er Objektiv"）。
出力:
  raw_heine/qpcr_long.csv      condition, cultivation, pool(total|viable), day, species, rep, pct
  raw_heine/livedead_long.csv  condition, cultivation, day, image, live_pct, dead_pct
  raw_heine/volume_long.csv    condition, cultivation, day, image, volume, objective_note
"""
import csv
from pathlib import Path

import openpyxl

D = Path(__file__).resolve().parents[1] / "data_5species" / "experiment_data" / "raw_heine"
SPN = {"S. oralis": "S. oralis", "A. naeslundii": "A. naeslundii", "V. dispar": "V. dispar",
       "V. parvula": "V. parvula", "F. nucleatum": "F. nucleatum",
       "P. gingivalis20709": "P. gingivalis_20709", "P. gingivalisW83": "P. gingivalis_W83"}


def qpcr():
    wb = openpyxl.load_workbook(D / "qpcr_species_distribution.xlsx", data_only=True)
    out = []
    for cond in ("Commensal", "Dysbiotic"):
        rows = list(wb[cond].iter_rows(values_only=True))
        for i, r in enumerate(rows):
            if not (isinstance(r[0], str) and r[0] != "Tag"):
                continue
            cult, _, kind = r[0].partition(" ")
            pool = "total" if kind.startswith("all") else "viable"
            hdr = rows[i + 1]
            st = [j for j, c in enumerate(hdr) if isinstance(c, str) and c != "Tag"] + [len(hdr)]
            names = [SPN[hdr[j].strip()] for j in st[:-1]]
            for d in rows[i + 2:]:
                if d[0] is None or isinstance(d[0], str):
                    break
                for k, sp in enumerate(names):
                    vals = [c for c in d[st[k]:st[k + 1]] if isinstance(c, (int, float))]
                    for rep, v in enumerate(vals, 1):
                        out.append([cond, cult, pool, int(d[0]), sp, rep, float(v)])
    return out


def livedead():
    wb = openpyxl.load_workbook(D / "livedead_HOBIC_Static.xlsx", data_only=True)
    sheets = {"HOBIC22": ("Commensal", "HOBIC"), "Static22": ("Commensal", "Static"),
              "HOBIC24": ("Dysbiotic", "HOBIC"), "Static24": ("Dysbiotic", "Static")}
    ld, vol = [], []
    for sh, (cond, cult) in sheets.items():
        rows = list(wb[sh].iter_rows(values_only=True))
        hdr = rows[1]
        vdays = [int(str(h).split()[1]) for h in hdr[:6]]
        conv = next((i for i, r in enumerate(rows) if r[0] == "Umgerechnet auf 63er Objektiv"), None)
        end = conv if conv is not None else len(rows)
        for img, r in enumerate(rows[2:end], 1):
            for j, day in enumerate(vdays):
                if isinstance(r[j], (int, float)):
                    note = "40x (see 63x-converted rows)" if (sh == "HOBIC22" and day == 1) else ""
                    vol.append([cond, cult, day, img, float(r[j]), note])
            for j in range(7, len(hdr), 2):
                if hdr[j] is None:
                    continue
                day = int(str(hdr[j]).split()[1])
                lv, dd = r[j], r[j + 1] if j + 1 < len(r) else None
                if isinstance(lv, (int, float)):
                    ld.append([cond, cult, day, img, float(lv), float(dd) if isinstance(dd, (int, float)) else ""])
        if conv is not None:
            for img, r in enumerate(rows[conv + 1:], 1):
                if isinstance(r[0], (int, float)):
                    vol.append([cond, cult, 1, img, float(r[0]), "63x-converted"])
    return ld, vol


def write(name, header, rows):
    with open(D / name, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    print(f"{name}: {len(rows)} rows")


if __name__ == "__main__":
    write("qpcr_long.csv", ["condition", "cultivation", "pool", "day", "species", "rep", "pct"], qpcr())
    ld, vol = livedead()
    write("livedead_long.csv", ["condition", "cultivation", "day", "image", "live_pct", "dead_pct"], ld)
    write("volume_long.csv", ["condition", "cultivation", "day", "image", "volume", "objective_note"], vol)
